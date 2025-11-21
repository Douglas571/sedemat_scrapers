from openpyxl import load_workbook
from pydantic import BaseModel

class AccountStatementData(BaseModel):
  initial_amount: float
  final_amount: float


def load_9290_account_statement_data(month: int, year: int) -> AccountStatementData:

  initial_amount = 0
  final_amount = 0

  return AccountStatementData(initial_amount=0, final_amount=0)

def load_1892_account_statement_data(month: int, year: int) -> AccountStatementData:

  initial_amount = 0
  final_amount = 0

  path = f'datos/account_statements/{str(year)[-2:]}-{month:02d}-1892.xlsx'
  wb = load_workbook(path)
  ws = wb['data']

  initial_amount = ws['D2'].value

  row = 2
  while True:
    cell = ws[f'D{row}']
    if cell.value is None:
      break
    final_amount = cell.value if cell.value is not None else 0
    row += 1

  initial_amount = float(initial_amount.replace('.', '').replace(',', '.'))
  final_amount = float(final_amount.replace('.', '').replace(',', '.'))

  return AccountStatementData(initial_amount=initial_amount, final_amount=final_amount)

def load_account_statement_data(month: int, year: int, bank_account: str) -> AccountStatementData:

  if bank_account == '1892':
    return load_1892_account_statement_data(month, year)

  if bank_account == '9290':
    return load_9290_account_statement_data(month, year)
  
if __name__ == "__main__":

  month = 5
  year = 2025
  bank_account = '1892'

  print(load_account_statement_data(month, year, bank_account))