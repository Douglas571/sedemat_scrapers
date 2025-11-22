from datetime import datetime
import json
from openpyxl import load_workbook

from project_types import Settlement

from settings import *

SETTLEMENTS_COLUMNS_MAPPER = {
  "legal_name": 0,
  "rif_cedula": 1,
  "code": 2,
  "description": 3,
  "paid_at_raw": 4,
  "settled_at": 5,
  "account_number": 6,
  "bank": 7,
  "reference_raw": 8,
  "amount": 9,
}


def load_non_standardized_settlements(path: str = DEFAULT_SETTLEMENTS_PATH) -> list[Settlement]:
  settlements_list = []
  workbook = load_workbook(path)
  sheet = workbook.active
  for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):

    if index == 1:
      continue

    settlement = {}
    settlement['legal_name'] = str(row[0]).strip()
    settlement['rif_cedula'] = str(row[1]).strip()
    settlement['num_comprobante'] = str(row[2]).strip()
    settlement['pago_por'] = str(row[3]).strip()
    settlement['fecha_pago'] = str(row[4]).strip()
    settlement['fecha'] = str(row[5]).strip()
    settlement['cuenta'] = str(row[6]).strip()
    settlement['banco'] = str(row[7]).strip()
    
    settlement['referencia'] = str(row[8]).strip()
    settlement['monto'] = row[9]
    settlement['payments'] = []

    settlement['not_found_payments'] = ''

    settlement['amount_difference'] = 0


    is_exonerated = any([
      'EXONERADO' in str(x).upper() if x is not None else False
      for x in [settlement['referencia'], settlement['monto'], settlement['banco'], settlement['cuenta']]
    ]) or settlement['referencia'] == "None"

    settlement['is_verified'] = True
    # settlement['comment'] = ''

    settled_at = settlement['fecha']
    if type(settled_at) == str:
      settled_at = datetime.strptime(settled_at, '%Y-%m-%d %H:%M:%S').date()

    reference = []
    paid_at = []
    not_found_payments = []
    bank=''
    account_number=''

    if not is_exonerated:
      references = str(settlement['referencia'])
      reference = references.replace("/", " ").replace("-", " ").split(" ")
      reference = [ref.strip() for ref in reference if ref.strip() != '']

      not_found_payments = reference.copy()

      bank = settlement['banco']
      account_number = settlement['cuenta']

      paid_at_raw = settlement['fecha_pago']
      if type(paid_at_raw) == str:
        formats = ["%d/%m/%Y", "%Y-%m-%d %H:%M:%S"]
        paid_at = []
        for fmt in formats:
          try:
            paid_at = [datetime.strptime(paid_at_raw, fmt).date()]
            break
          except ValueError:
            pass
            # print(f"Warning: Could not parse paid_at date '{paid_at_raw}' with format {fmt}")
            
        if not paid_at:
          if SHOW_WARNINGS: print(f"Warning: Could not parse paid_at date '{paid_at_raw}' for settlement {settlement['num_comprobante']}")

      else:
        paid_at = [paid_at_raw]

    

    std_settlement = Settlement(
      legal_name = settlement['legal_name'],
      rif_cedula = settlement['rif_cedula'],
      code = settlement['num_comprobante'],
      description = settlement['pago_por'],
      settled_at = settled_at,

      account_number = account_number,
      bank = bank,

      amount = float(settlement['monto']) if not is_exonerated else 0.0,
      reference = reference,
      reference_raw = settlement['referencia'],

      paid_at = paid_at,
      paid_at_raw = str(settlement['fecha_pago']),

      not_found_payments=not_found_payments,

      is_exonerated=is_exonerated,
    )
      
    settlements_list.append(std_settlement)
  return settlements_list

if __name__ == '__main__':
    settlements_list = load_non_standardized_settlements(
        './datos/settlements/cuadro_to_use.xlsx'
    )

    for settlement in settlements_list:
        # print(json.dumps([settlement.model_dump()], indent=2, default=str))
        print(list(settlement.model_dump().keys()))
