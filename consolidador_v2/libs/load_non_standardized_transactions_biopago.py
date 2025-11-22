import json
from openpyxl import load_workbook
from datetime import datetime

from project_types import Transaction


def load_non_standardized_transactions_biopago(path: str) -> list[Transaction]:
  file_name = path

  try:
    workbook = load_workbook(file_name)
  except FileNotFoundError:
    print(f"Warning: file {file_name} not found")
    return []

  sheet = workbook.active
  transactions_list = []

  for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):

    if index == 1:
      continue

    # omit empty rows (those with amount = None)
    if row[4] == None: 
      continue

    # print(row)
    transaction = Transaction(
      date = datetime.strptime(str(row[1]), "%Y-%m-%d %H:%M:%S").date(),
      reference=str(row[0]).strip(),
      description=f"{row[7]} {row[5]}",
      amount=float(row[4]),
      bank="BIOPAGO",
      account_number="1892"
    )

    transactions_list.append(transaction)

  return transactions_list

if __name__ == "__main__":
  path = 'datos/account_statements/25-08-biopago.xlsx'
  transactions_list = load_non_standardized_transactions_biopago(path)
  transactions_list.sort(key=lambda t: t.date)
  print(json.dumps([t.model_dump() for t in transactions_list], indent=2, default=str))
